# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

# 1. إنشاء كتاب عمل جديد
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sorbitol Liquid Assay"

# 2. إعدادات التنسيق والألوان الرسمية للمختبرات
thin = Side(style='thin', color='000000')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
title_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
header_fill = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')
white_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

# ضبط أبعاد الأعمدة لسهولة القراءة
widths = {'A': 25, 'B': 18, 'C': 18, 'D': 18}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

# 3. العنوان الرئيسي
ws.merge_cells('A1:D1')
ws['A1'] = "Sorbitol Liquid Assay Report (As Is Basis)"
ws['A1'].font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws['A1'].fill = title_fill
ws.row_dimensions[1].height = 30

# 4. بيانات العينة الأساسية
ws['A3'] = "Batch Number:"
ws['B3'] = "26100170"
ws['A4'] = "Instrument / Method:"
ws['B4'] = "HPLC / HPMC"
for r in [3, 4]:
    ws[f'A{r}'].font = Font(name='Calibri', bold=True)
    ws[f'B{r}'].font = Font(name='Calibri')

# 5. جدول حقنات المحلول القياسي (Standard)
ws['A6'] = "Standard Injections"
ws['A6'].font = Font(name='Calibri', bold=True, size=11)
ws.merge_cells('A7:B7')
ws['A7'] = "Injection No."
ws['C7'] = "Peak Area"
for col in ['A', 'B', 'C']:
    ws[f'{col}7'].font = Font(name='Calibri', bold=True)
    ws[f'{col}7'].fill = header_fill
    ws[f'{col}7'].border = box
    ws[f'{col}7'].alignment = Alignment(horizontal='center')

std_data = [("Standard Inj 1", 6585044), ("Standard Inj 2", 6528789), ("Standard Inj 3", 6554353)]
for i, (name, area) in enumerate(std_data, start=8):
    ws.merge_cells(f'A{i}:B{i}')
    ws[f'A{i}'] = name
    ws[f'C{i}'] = area
    ws[f'A{i}'].border = box
    ws[f'C{i}'].border = box
    ws[f'C{i}'].number_format = '#,##0'

# حسابات القياسي المدمجة بالمعادلات
ws.merge_cells('A11:B11')
ws['A11'] = "Standard Average Area:"
ws['C11'] = "=AVERAGE(C8:C10)"
ws['A12'] = "RSD (%)"
ws['C12'] = "=(STDEV(C8:C10)/C11)*100"

# 6. جدول حقنات العينة السائلة (Sample As Is)
ws['A14'] = "Sample Injections (Sorbitol Liquid)"
ws['A14'].font = Font(name='Calibri', bold=True, size=11)
ws.merge_cells('A15:B15')
ws['A15'] = "Injection No."
ws['C15'] = "Peak Area"
for col in ['A', 'B', 'C']:
    ws[f'{col}15'].font = Font(name='Calibri', bold=True)
    ws[f'{col}15'].fill = header_fill
    ws[f'{col}15'].border = box
    ws[f'{col}15'].alignment = Alignment(horizontal='center')

sam_data = [("Sample Inj 4", 6553328), ("Sample Inj 5", 6559900), ("Sample Inj 6", 6595998)]
for i, (name, area) in enumerate(sam_data, start=16):
    ws.merge_cells(f'A{i}:B{i}')
    ws[f'A{i}'] = name
    ws[f'C{i}'] = area
    ws[f'A{i}'].border = box
    ws[f'C{i}'].border = box
    ws[f'C{i}'].number_format = '#,##0'

ws.merge_cells('A19:B19')
ws['A19'] = "Sample Average Area:"
ws['C19'] = "=AVERAGE(C16:C18)"

# تطبيق التنسيقات النهائية للجداول وحساب النتيجة الإجمالية المباشرة للسائل
for r in [11, 12, 19]:
    ws[f'A{r}'].font = Font(name='Calibri', bold=True)
    ws[f'C{r}'].font = Font(name='Calibri', bold=True)
    ws[f'A{r}'].border = box
    ws[f'C{r}'].border = box

ws['C11'].number_format = '#,##0'
ws['C19'].number_format = '#,##0'
ws['C12'].number_format = '0.00"%"'

# 7. قسم الحسابات التحليلية النهائية للأوزان والـ Assay المباشر
ws['A21'] = "Variables & Parameters"
ws['A21'].font = Font(name='Calibri', bold=True, size=11)

params = [
    ("Standard Weight (g)", 1.0000),
    ("Sample Weight (g)", 1.0000),
    ("Standard Dilution (ml)", 100),
    ("Sample Dilution (ml)", 100),
    ("Standard Purity (%)", 100.0)
]
for i, (p_name, p_val) in enumerate(params, start=22):
    ws[f'A{i}'] = p_name
    ws[f'B{i}'] = p_val
    ws[f'A{i}'].font = Font(name='Calibri')
    ws[f'B{i}'].font = Font(name='Calibri')
    ws[f'A{i}'].border = box
    ws[f'B{i}'].border = box

# المعادلة الرياضية النهائية لحساب النسبة في السائل (As Is) مباشرة
ws['A28'] = "Final Assay (As Is) (%)"
ws['B28'] = "=(C19/C11)*(B22/B23)*(B25/B24)*(B26/100)*100"
ws['A28'].font = Font(name='Calibri', bold=True, color='1F4E78')
ws['B28'].font = Font(name='Calibri', bold=True, color='1F4E78')
ws['A28'].border = box
ws['B28'].border = box
ws['B28'].number_format = '0.00"%"'

# حفظ ملف الإكسيل النهائي في مجلد عملك بالويندوز
wb.save(r'D:\Projects\PayRolls\Sorbitol_Liquid_Assay_Report.xlsx')
print("تم توليد ملف الإكسيل بنجاح في المسار المحدد!")
